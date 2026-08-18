#!/usr/bin/env python3
"""
categorize_expenses.py  —  extensible, fail-loud monthly expense aggregator
===========================================================================
Turns a bank statement (+ optional credit-card bill PDFs) into a REVIEWABLE
draft whose "Entries draft" tab pastes straight into Monthly Budget Tracker.xlsx.

TWO DESIGN PRINCIPLES
---------------------
1. TEMPLATE / PLUGGABLE FORMATS.  Bank layouts live in BANK_FORMATS and card
   issuers in CARD_FORMATS. Each is a small, self-contained entry. The UOB
   entries are worked examples — copy one, edit it, and a new bank or card is
   supported without touching the engine.

2. NO SILENT FAILURES.  Every step validates and reconciles. If a format is not
   recognised, a bank row will not parse, a rules row has an unknown Type
   (Income/Actual/One-off) or Tier (Core/Over & above/One-off/Reimbursable/
   Investment), a card bill's line items do not reconcile to its printed total,
   an attached statement has no readable total, or no statement total matches the
   card payment it should pay, the tool records a BLOCKING issue, keeps the
   affected money visible (a card payment stays a lump rather than vanishing),
   writes a "Validation" tab, prints a summary, and exits non-zero.

   Bookings honour the rule's own Type and Tier: an include=N line is skipped; a
   merchant refund reduces spending; a card payment is never counted as spending;
   a deposit that matches an expense rule (e.g. a reimbursement) is booked as
   NEGATIVE spending, and an unclassified deposit is surfaced for review rather
   than dropped. Each card payment is matched to EXACTLY ONE statement by amount
   (never by shared last-four digits), so two statements for the same card cannot
   double-count. It never guesses and never drops money quietly.

USAGE
-----
    python3 categorize_expenses.py BANK.(xlsx|csv) --month 2026-06 \
        --bill UOB.pdf --bill OCBC.pdf --rules rules.csv --out "Jun2026 draft.xlsx"

Exit codes:  0 = PASS or REVIEW (open the draft, resolve amber flags)
             2 = FAIL (blocking issues — see the Validation tab; do NOT paste yet)

--------------------------------------------------------------------------
HOW TO ADD A NEW BANK  (see the TEMPLATE entry in BANK_FORMATS):
  * 'detect' is a list of (column_index, text) the header row must contain IN
    THAT COLUMN — this both identifies the format and verifies column positions,
    so a reordered/renamed export fails detection instead of being misread.
  * 'columns' maps fields to columns and states the amount convention.
HOW TO ADD A NEW CARD  (see parse_uob_card + the TEMPLATE in CARD_FORMATS):
  * 'detect(text)' returns True for that issuer's PDF.
  * 'parse(text)' returns dict(last4, items[{desc,amount,credit}],
    previous_balance, total_balance). The engine reconciles
    previous_balance + sum(signed items) == total_balance for you.
--------------------------------------------------------------------------
"""
import argparse, csv, re, sys, datetime
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

TOL = 0.01  # reconciliation tolerance ($)

# =====================================================================
#  BANK FORMAT REGISTRY   (add a dict here to support a new bank)
# =====================================================================
BANK_FORMATS = [
    {
        'name': 'UOB One Account',
        'source': ('xlsx', 'csv'),          # accepted file types
        # header row must contain these tokens IN these 1-based columns:
        'detect': [(1, 'Transaction Date'), (3, 'Withdrawal'), (4, 'Deposit')],
        'columns': {
            'date': 1,
            'desc': [2],                     # one or more columns, joined
            'mode': 'wd_dp',                 # separate Withdrawal / Deposit cols
            'withdrawal': 3,
            'deposit': 4,
        },
        'date_formats': ['%d %b %Y', '%d/%m/%Y', '%Y-%m-%d'],
    },
    # ------------------------------------------------------------------
    # TEMPLATE — copy, rename, and edit for a new bank. Delete if unused.
    # {
    #     'name': 'Example Bank (single signed Amount column)',
    #     'source': ('csv', 'xlsx'),
    #     'detect': [(1, 'Date'), (2, 'Description'), (3, 'Amount')],
    #     'columns': {
    #         'date': 1, 'desc': [2],
    #         'mode': 'signed', 'amount': 3,
    #         'debit_sign': 'negative',   # 'negative' = spend is negative; else 'positive'
    #     },
    #     'date_formats': ['%d/%m/%Y', '%m/%d/%Y'],
    # },
]

# =====================================================================
#  CARD BILL REGISTRY   (add a parser here to support a new card issuer)
# =====================================================================
CARD_RE = re.compile(r'(\d{4})[- ]?(\d{4})[- ]?(\d{4})[- ]?(\d{4})')
def card_last4(text):
    m = CARD_RE.search(text.replace(' ', ''))
    return m.group(0)[-4:] if m else None

_UOB_LINE = re.compile(r'^(\d{2} [A-Z]{3})\s+(\d{2} [A-Z]{3})\s+(.*?)\s+([\d,]+\.\d{2})(CR)?$')
_MONEY = re.compile(r'([\d,]+\.\d{2})')
def parse_uob_card(text):
    """UOB card statement -> dict(last4, items, previous_balance, total_balance)."""
    last4 = card_last4(text)
    prev = total = None
    items = []
    for ln in text.split('\n'):
        s = ln.strip()
        up = s.upper()
        if up.startswith('PREVIOUS BALANCE'):
            m = _MONEY.search(s);  prev = float(m.group(1).replace(',', '')) if m else prev
            continue
        if up.startswith('TOTAL BALANCE') or up.startswith('SUB TOTAL'):
            m = _MONEY.search(s);  total = float(m.group(1).replace(',', '')) if m else total
            continue
        m = _UOB_LINE.match(s)
        if not m:
            continue
        desc = m.group(3).strip()
        amt = float(m.group(4).replace(',', ''))
        credit = bool(m.group(5)) or 'REBATE' in desc.upper() or 'PAYMT' in desc.upper()
        items.append({'desc': desc, 'amount': amt, 'credit': credit})
    # UOB prints SUB TOTAL = new balance after the prev-balance payment cleared it;
    # normalise: if we only saw SUB TOTAL, previous balance is effectively settled.
    if prev is None:
        prev = 0.0
    return {'last4': last4, 'items': items, 'previous_balance': prev, 'total_balance': total}

CARD_FORMATS = [
    {
        'name': 'UOB Card',
        'detect': lambda t: ('UNITED OVERSEAS BANK' in t.upper() or 'UOB' in t.upper())
                            and card_last4(t) is not None,
        'parse': parse_uob_card,
    },
    # ------------------------------------------------------------------
    # TEMPLATE — add another issuer. 'parse' must return the same dict shape;
    # the engine reconciles previous_balance + sum(signed items) == total_balance.
    # {
    #     'name': 'OCBC Card',
    #     'detect': lambda t: 'OVERSEA-CHINESE BANKING' in t.upper() or 'OCBC' in t.upper(),
    #     'parse': parse_ocbc_card,
    # },
]

# =====================================================================
#  helpers: load a grid, parse values
# =====================================================================
def load_grid(path):
    if path.lower().endswith('.csv'):
        with open(path, newline='') as f:
            return [list(r) for r in csv.reader(f)]
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]

def cell(grid, r, c):
    return grid[r][c - 1] if 0 <= r < len(grid) and 0 < c <= len(grid[r]) else None

def to_float(x):
    if x is None or x == '':
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).replace(',', '').replace('$', '').strip()
    try:
        return float(s)
    except ValueError:
        return None

def to_date(x, fmts):
    if isinstance(x, datetime.datetime):
        return x
    if isinstance(x, datetime.date):
        return datetime.datetime(x.year, x.month, x.day)
    if x is None:
        return None
    s = str(x).strip()
    for f in fmts:
        try:
            return datetime.datetime.strptime(s, f)
        except ValueError:
            continue
    return None

# =====================================================================
#  bank parsing with detection + row-level validation
# =====================================================================
def detect_bank_format(grid, ftype):
    for fmt in BANK_FORMATS:
        if ftype not in fmt['source']:
            continue
        for r in range(min(len(grid), 40)):
            ok = True
            for col, token in fmt['detect']:
                v = cell(grid, r, col)
                if v is None or token.lower() not in str(v).lower():
                    ok = False
                    break
            if ok:
                return fmt, r
    return None, None

def parse_bank(path, issues):
    ftype = 'csv' if path.lower().endswith('.csv') else 'xlsx'
    grid = load_grid(path)
    fmt, hdr = detect_bank_format(grid, ftype)
    if fmt is None:
        supported = ', '.join(f"{f['name']}" for f in BANK_FORMATS)
        issues.append(('FAIL', 'bank-format',
            f'Bank file "{path}" did not match any known format (checked column positions of headers). '
            f'Supported: {supported}. Add a BANK_FORMATS entry for this layout.'))
        return None, []
    cols = fmt['columns']
    rows = []
    for r in range(hdr + 1, len(grid)):
        raw = grid[r]
        if raw is None or all(v is None or str(v).strip() == '' for v in raw):
            continue  # genuinely blank line
        dt = to_date(cell(grid, r, cols['date']), fmt['date_formats'])
        desc = ' | '.join(str(cell(grid, r, c)).replace('\n', ' | ')
                          for c in cols['desc'] if cell(grid, r, c) is not None).strip()
        if cols['mode'] == 'wd_dp':
            wd = to_float(cell(grid, r, cols['withdrawal'])) or 0.0
            dp = to_float(cell(grid, r, cols['deposit'])) or 0.0
            wd_raw = cell(grid, r, cols['withdrawal']); dp_raw = cell(grid, r, cols['deposit'])
            amt_ok = (to_float(wd_raw) is not None) or (to_float(dp_raw) is not None) or (wd_raw in (None,'',0) and dp_raw in (None,'',0))
            both = wd > 0 and dp > 0
        else:  # signed
            a = to_float(cell(grid, r, cols['amount']))
            amt_ok = a is not None
            both = False
            if a is None:
                wd = dp = 0.0
            elif cols.get('debit_sign', 'negative') == 'negative':
                wd, dp = (-a, 0.0) if a < 0 else (0.0, a)
            else:
                wd, dp = (a, 0.0) if a > 0 else (0.0, -a)
        # validation — surface, never skip
        if dt is None or not amt_ok or both or desc == '':
            problem = []
            if dt is None: problem.append('unparseable date')
            if not amt_ok: problem.append('non-numeric amount')
            if both: problem.append('both withdrawal & deposit set')
            if desc == '': problem.append('empty description')
            issues.append(('FAIL', 'bank-row',
                f'Row {r+1} in "{path}" could not be parsed ({", ".join(problem)}): {raw}'))
            continue
        rows.append({'row': r + 1, 'date': dt, 'desc': desc, 'out': wd, 'in': dp})
    issues.append(('INFO', 'bank-format', f'Detected bank format: {fmt["name"]} ({len(rows)} rows parsed).'))
    return fmt, rows

# =====================================================================
#  card bill parsing + reconciliation
# =====================================================================
def parse_card_bill(path, issues):
    if pdfplumber is None:
        issues.append(('FAIL', 'deps', 'pdfplumber not installed: pip3 install pdfplumber'))
        return None
    try:
        with pdfplumber.open(path) as pdf:
            text = '\n'.join((p.extract_text() or '') for p in pdf.pages)
    except Exception as e:
        issues.append(('FAIL', 'bill-open', f'Could not open bill "{path}": {e}'))
        return None
    fmt = next((f for f in CARD_FORMATS if f['detect'](text)), None)
    if fmt is None:
        issues.append(('FAIL', 'bill-format',
            f'Card bill "{path}" did not match any known issuer. '
            f'Supported: {", ".join(f["name"] for f in CARD_FORMATS)}. Add a CARD_FORMATS parser.'))
        return None
    d = fmt['parse'](text)
    d['name'] = fmt['name']; d['path'] = path
    # RECONCILE: previous_balance + sum(signed items) == total_balance
    signed = sum((-it['amount'] if it['credit'] else it['amount']) for it in d['items'])
    d['charges'] = round(sum(it['amount'] for it in d['items'] if not it['credit']), 2)
    d['credits'] = round(sum(it['amount'] for it in d['items'] if it['credit']), 2)
    if not d['items']:
        issues.append(('FAIL', 'bill-empty',
            f'Bill "{path}" ({fmt["name"]}, card …{d.get("last4")}) yielded NO line items — '
            f'parser likely out of date. The matching payment will stay a lump (not silently dropped).'))
        d['reconciled'] = False
    elif d['total_balance'] is None:
        issues.append(('REVIEW', 'bill-total',
            f'Bill "{path}" has no readable total to reconcile against; using {len(d["items"])} parsed items — please spot-check.'))
        d['reconciled'] = True
    elif abs(d['previous_balance'] + signed - d['total_balance']) > TOL:
        issues.append(('FAIL', 'bill-reconcile',
            f'Bill "{path}" (card …{d.get("last4")}) does NOT reconcile: previous {d["previous_balance"]:.2f} '
            f'+ lines {signed:.2f} = {d["previous_balance"]+signed:.2f} vs printed total {d["total_balance"]:.2f}. '
            f'Some lines were missed — payment kept as a lump, not silently broken down.'))
        d['reconciled'] = False
    else:
        d['reconciled'] = True
        issues.append(('INFO', 'bill',
            f'Bill {fmt["name"]} card …{d["last4"]}: {len(d["items"])} lines, charges {d["charges"]:.2f}, reconciled OK.'))
    return d

# =====================================================================
#  rules
# =====================================================================
def load_rules(path, issues):
    rules = []
    try:
        with open(path, newline='') as f:
            for row in csv.reader(f):
                if not row or row[0].strip().startswith('#') or row[0].strip() == 'pattern':
                    continue
                pat = row[0].strip()
                if not pat:
                    continue
                try:
                    rx = re.compile(pat, re.I)
                except re.error as e:
                    issues.append(('REVIEW', 'rules', f'Bad regex in rules.csv: "{pat}" ({e}) — skipped.'))
                    continue
                raw_typ = (row[2].strip() if len(row) > 2 else '')
                typ = {'': 'Actual', 'income': 'Income', 'actual': 'Actual',
                       'one-off': 'One-off', 'oneoff': 'One-off', 'one off': 'One-off'}.get(raw_typ.lower())
                if typ is None:
                    issues.append(('FAIL', 'rules',
                        f'Rule "{pat}" has an invalid Type "{raw_typ}" — must be Income, Actual or One-off. '
                        f'Fix rules.csv; this rule was skipped so nothing is silently mis-booked or dropped.'))
                    continue
                raw_tier = (row[5].strip() if len(row) > 5 else '')
                if raw_tier == '':
                    tier = ''   # classify() defaults it (— for Income, else Core)
                else:
                    tier = {'core': 'Core', 'over & above': 'Over & above', 'over and above': 'Over & above',
                            'over&above': 'Over & above', 'one-off': 'One-off', 'oneoff': 'One-off',
                            'one off': 'One-off', 'reimbursable': 'Reimbursable', 'reimburse': 'Reimbursable',
                            'investment': 'Investment', '—': '—', '-': '—'}.get(raw_tier.lower())
                    if tier is None:
                        issues.append(('FAIL', 'rules',
                            f'Rule "{pat}" has an invalid Tier "{raw_tier}" — must be Core, Over & above, One-off, '
                            f'Reimbursable or Investment. Fix rules.csv; this rule was skipped so nothing lands in '
                            f'the wrong tier or vanishes from the lifestyle metrics.'))
                        continue
                rules.append((rx,
                              (row[1].strip() if len(row) > 1 else '') or '',
                              typ,
                              (row[3].strip().upper() if len(row) > 3 else 'Y') == 'Y',
                              row[4].strip() if len(row) > 4 else '',
                              tier))
    except FileNotFoundError:
        issues.append(('FAIL', 'rules', f'Rules file not found: {path}'))
    return rules

def classify(desc, rules):
    for rx, cat, typ, inc, note, tier in rules:
        if rx.search(desc):
            if not tier: tier = '—' if typ == 'Income' else 'Core'
            return (cat or '(exclude)'), typ, inc, note, tier
    return 'Uncategorised', 'Actual', True, 'REVIEW - no rule matched', 'Core'

# =====================================================================
#  draft workbook
# =====================================================================
A = 'Arial'
def _f(sz=10, b=False, color='000000'):
    return Font(name=A, size=sz, bold=b, color=color)
HF = Font(name=A, size=10, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F4E5F')
FLAG = PatternFill('solid', fgColor='FCE4A6'); EXC = PatternFill('solid', fgColor='EEEEEE')
FAILF = PatternFill('solid', fgColor='F4CCCC'); OKF = PatternFill('solid', fgColor='D9EAD3')
INFOF = PatternFill('solid', fgColor='E7F0F7')
M2 = '#,##0.00'
def _hdr(ws, cols):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = HF; c.fill = HFILL
        c.alignment = Alignment('center', wrap_text=True)

def build_draft(bank_rows, bills, month, rules, issues, out_path):
    agg = defaultdict(float)
    flagged = 0
    recon = {'income_in': 0.0, 'actual_bank': 0.0, 'actual_card': 0.0, 'card_lump': 0.0}
    # keep EVERY attached statement (two statements can share a card number); each is
    # matched to at most one payment, by amount, and only the matched one is broken down.
    bills_by4 = defaultdict(list)
    for bb in bills:
        if bb and bb.get('last4'):
            bills_by4[bb['last4']].append(bb)
    seen_cards = set()      # last4 of card payments seen in the bank statement
    consumed = set()        # id() of statements matched to a payment (no input mutation)

    # month-scope sanity: every draft entry is dated to `month`, so warn if the bank
    # rows actually span other months (import one statement month at a time).
    row_months = {(t['date'].year, t['date'].month) for t in bank_rows
                  if isinstance(t.get('date'), datetime.datetime)}
    off_months = sorted(m for m in row_months if m != (month.year, month.month))
    if off_months:
        issues.append(('REVIEW', 'month-scope',
            f'Bank rows include date(s) outside {month.strftime("%Y-%m")} '
            f'({", ".join("%04d-%02d" % m for m in off_months)}); every draft entry is dated '
            f'{month.strftime("%Y-%m-01")}. Import one statement month at a time for clean monthly totals.'))

    out = openpyxl.Workbook()

    # ---- Bank tab ----
    b = out.active; b.title = 'Bank'
    _hdr(b, ['Row', 'Date', 'Transaction', 'Out', 'In', 'Category', 'Type', 'Include?', 'Flag / Note', 'Tier'])
    rr = 2
    for tx in bank_rows:
        cat, typ, inc, note, tier = classify(tx['desc'], rules)
        l4 = card_last4(tx['desc'])
        is_card = bool(l4) and re.search(r'card|CC|bill payment', tx['desc'], re.I) and tx['out'] > 0
        b.cell(row=rr, column=1, value=tx['row']).font = _f(9, color='999999')
        b.cell(row=rr, column=2, value=tx['date']).number_format = 'yyyy-mm-dd'
        b.cell(row=rr, column=3, value=tx['desc']).font = _f(9)
        if tx['out']: b.cell(row=rr, column=4, value=tx['out']).number_format = M2
        if tx['in']:  b.cell(row=rr, column=5, value=tx['in']).number_format = M2
        b.cell(row=rr, column=6, value=('—' if not inc else cat)).font = _f()
        b.cell(row=rr, column=7, value=typ).font = _f()
        b.cell(row=rr, column=8, value=('Yes' if inc else 'No')).font = _f()
        b.cell(row=rr, column=9, value=note).font = _f(9)
        b.cell(row=rr, column=10, value=('' if not inc else tier)).font = _f()
        if is_card:
            seen_cards.add(l4)
            cands = [bb for bb in bills_by4.get(l4, []) if bb.get('reconciled') and id(bb) not in consumed]
            # match this payment to exactly ONE unconsumed statement whose readable printed
            # total equals the payment amount — never by shared card digits alone.
            match = next((bb for bb in cands
                          if bb.get('total_balance') is not None
                          and abs(tx['out'] - bb['total_balance']) <= TOL), None)
            if match is not None:
                consumed.add(id(match))
                b.cell(row=rr, column=8, value='Replaced')
                b.cell(row=rr, column=9, value=f'Broken down from reconciled bill (card …{l4}, total {match["total_balance"]:.2f})')
                b.cell(row=rr, column=10, value='')
            else:
                # keep as lump — NEVER silently drop; explain why it was not broken down
                b.cell(row=rr, column=6, value='Credit Card'); b.cell(row=rr, column=10, value='Core')
                agg[('Credit Card', 'Actual', 'Core')] += tx['out']; recon['card_lump'] += tx['out']
                allb = bills_by4.get(l4, [])
                if not allb:
                    b.cell(row=rr, column=9, value='Kept as lump — no bill attached (normal)')
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = INFOF
                elif not any(bb.get('reconciled') for bb in allb):
                    b.cell(row=rr, column=9, value='REVIEW - bill did NOT reconcile; kept as lump (see Validation tab)')
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
                elif any(bb.get('reconciled') and bb.get('total_balance') is None for bb in allb):
                    b.cell(row=rr, column=9, value='REVIEW - bill total unreadable, cannot verify against payment; '
                           'kept as lump (fix the parser or the bill)')
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
                else:
                    b.cell(row=rr, column=9, value=f'REVIEW - no attached statement total matches payment '
                           f'{tx["out"]:.2f}; kept as lump (right month?)')
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
        elif not inc:
            for cc in range(1, 11): b.cell(row=rr, column=cc).fill = EXC
            if 'REVIEW' in note: flagged += 1
        else:
            if tx['out'] > 0:
                if 'REVIEW' in note:
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
                    flagged += 1
                otyp = typ if typ in ('Actual', 'One-off') else 'Actual'   # honour the rule's type (e.g. One-off investment); a mis-typed Income on an outflow falls back to Actual
                agg[(cat, otyp, tier)] += tx['out']; recon['actual_bank'] += tx['out']
            elif tx['in'] > 0:
                if typ == 'Income':
                    if 'REVIEW' in note:
                        for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
                        flagged += 1
                    agg[(cat, 'Income', '—')] += tx['in']; recon['income_in'] += tx['in']
                elif 'no rule matched' in note:
                    # unclassified deposit — do NOT guess its direction; surface it loudly
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = FLAG
                    flagged += 1
                    issues.append(('REVIEW', 'deposit',
                        f'Row {tx["row"]}: unclassified deposit {tx["in"]:.2f} "{tx["desc"][:40]}" — add a rule '
                        f'(Income, or an expense category if it is a refund/reimbursement). Left out of the draft.'))
                else:
                    # deposit that matches an EXPENSE rule = money back (refund / reimbursement)
                    # -> negative spending in that category, so net-of-reimbursable is right.
                    otyp = typ if typ in ('Actual', 'One-off') else 'Actual'
                    agg[(cat, otyp, tier)] -= tx['in']; recon['actual_bank'] -= tx['in']
                    b.cell(row=rr, column=9, value=((note + ' | ') if note else '')
                           + 'deposit booked as negative spending (refund/reimbursement)')
                    for cc in range(1, 11): b.cell(row=rr, column=cc).fill = INFOF
        rr += 1
    for col, w in zip('ABCDEFGHIJ', [6, 11, 44, 10, 10, 15, 8, 9, 40, 13]): b.column_dimensions[col].width = w
    b.freeze_panes = 'B2'

    # ---- one tab per reconciled bill (drill-down) ----
    for bill in bills:
        if not bill: continue
        cbs = out.create_sheet(f'Card …{bill.get("last4","?")}'[:31])
        _hdr(cbs, ['Merchant / Description', 'Amount', 'Category', 'Tier', 'Flag / Note'])
        rr = 2
        for it in bill['items']:
            matched = id(bill) in consumed   # this exact statement was matched to a payment
            if it['credit']:
                # a credit is either a card PAYMENT (excluded — it is the bill payment
                # we already handle on the bank side) or a merchant REFUND/REBATE, which
                # must REDUCE spending in its category.
                is_payment = bool(re.search(r'PAYMT|PAYMENT|AUTOPAY|GIRO\s*(DED|PYT|PAYMENT)', it['desc'], re.I))
                cbs.cell(row=rr, column=1, value=it['desc']).font = _f(9)
                cbs.cell(row=rr, column=2, value=-it['amount']).number_format = M2
                if is_payment:
                    cbs.cell(row=rr, column=3, value='(payment — excluded)').font = _f(9)
                    for cc in range(1, 6): cbs.cell(row=rr, column=cc).fill = EXC
                else:
                    cat, typ, inc, note, tier = classify(it['desc'], rules)
                    otyp = typ if typ in ('Actual', 'One-off') else 'Actual'
                    cbs.cell(row=rr, column=3, value=('(excluded by rule)' if not inc else f'{cat} (refund)')).font = _f()
                    cbs.cell(row=rr, column=4, value=('' if not inc else tier)).font = _f()
                    cbs.cell(row=rr, column=5, value=note).font = _f(9)
                    for cc in range(1, 6): cbs.cell(row=rr, column=cc).fill = INFOF
                    if matched and inc:
                        agg[(cat, otyp, tier)] -= it['amount']; recon['actual_card'] -= it['amount']
            else:
                # a bill's charges are only added when it reconciled AND was matched to a
                # real card payment of the right amount (matched, computed above)
                cat, typ, inc, note, tier = classify(it['desc'], rules)
                otyp = typ if typ in ('Actual', 'One-off') else 'Actual'   # a card charge is an outflow
                cbs.cell(row=rr, column=1, value=it['desc']).font = _f(9)
                cbs.cell(row=rr, column=2, value=it['amount']).number_format = M2
                cbs.cell(row=rr, column=3, value=('(excluded by rule)' if not inc else cat)).font = _f()
                cbs.cell(row=rr, column=4, value=('' if not inc else tier)).font = _f()
                cbs.cell(row=rr, column=5, value=note).font = _f(9)
                if not inc:
                    for cc in range(1, 6): cbs.cell(row=rr, column=cc).fill = EXC
                elif 'REVIEW' in note:
                    for cc in range(1, 6): cbs.cell(row=rr, column=cc).fill = FLAG
                    if matched: flagged += 1
                # honour include=N (skip) and the rule's type; never book an unmatched bill
                if matched and inc:
                    agg[(cat, otyp, tier)] += it['amount']; recon['actual_card'] += it['amount']
            rr += 1
        if not bill.get('reconciled'):
            st = 'NOT USED (failed reconciliation)'; fill = FAILF
        elif id(bill) in consumed:
            st = 'BROKEN DOWN (matched a card payment)'; fill = OKF
        else:
            st = 'RECONCILED but NOT matched to a payment — kept as lump'; fill = FLAG
        c = cbs.cell(row=rr + 1, column=1, value=f'Status: {st}   charges {bill.get("charges",0):.2f}, '
                     f'rebates/credits {bill.get("credits",0):.2f}, printed total {bill.get("total_balance")}')
        c.font = _f(9, b=True); c.fill = fill
        for col, w in zip('ABCDE', [46, 12, 18, 14, 40]): cbs.column_dimensions[col].width = w
        cbs.freeze_panes = 'A2'

    # ---- Entries draft ----
    ed = out.create_sheet('Entries draft')
    _hdr(ed, ['Date', 'Category', 'Type', 'Tier', 'Amount', 'Note'])
    order = {'Income': 0, 'Actual': 1, 'One-off': 2}
    rr = 2
    for (cat, typ, tier), amt in sorted(agg.items(), key=lambda kv: (order.get(kv[0][1], 9), kv[0][0], kv[0][2])):
        if abs(amt) < 0.005: continue
        ed.cell(row=rr, column=1, value=month).number_format = 'yyyy-mm-dd'
        ed.cell(row=rr, column=2, value=cat).font = _f()
        ed.cell(row=rr, column=3, value=typ).font = _f()
        ed.cell(row=rr, column=4, value=tier).font = _f()
        ed.cell(row=rr, column=5, value=round(amt, 2)).number_format = M2
        rr += 1
    ed.cell(row=rr + 1, column=2, value='Total Actual').font = _f(10, b=True)
    ed.cell(row=rr + 1, column=5, value=f'=SUMIF(C2:C{rr-1},"Actual",E2:E{rr-1})').number_format = M2
    ed.cell(row=rr + 2, column=2, value='Total One-off').font = _f(10, b=True)
    ed.cell(row=rr + 2, column=5, value=f'=SUMIF(C2:C{rr-1},"One-off",E2:E{rr-1})').number_format = M2
    ed.cell(row=rr + 3, column=2, value='Total Income').font = _f(10, b=True)
    ed.cell(row=rr + 3, column=5, value=f'=SUMIF(C2:C{rr-1},"Income",E2:E{rr-1})').number_format = M2
    for col, w in zip('ABCDEF', [12, 20, 10, 14, 13, 30]): ed.column_dimensions[col].width = w
    ed.freeze_panes = 'A2'

    # ---- attached-but-unmatched bills -> issue (one per statement, by identity) ----
    for bill in bills:
        if not bill or not bill.get('reconciled') or id(bill) in consumed:
            continue
        b4 = bill.get('last4')
        if bill.get('total_balance') is None:
            issues.append(('FAIL', 'unmatched-bill',
                f'Bill card …{b4} has no readable printed total, so it could not be verified against a card '
                f'payment — its charges were NOT added (kept as a lump). Fix the parser/bill or remove it.'))
        elif b4 in seen_cards:
            issues.append(('REVIEW', 'unmatched-bill',
                f'Bill card …{b4} (total {bill.get("total_balance"):.2f}) did not match any card payment amount — '
                f'kept as a lump (did you attach the right month, or is there a second statement for this card?).'))
        else:
            issues.append(('FAIL', 'unmatched-bill',
                f'Bill card …{b4} reconciled but no card payment for that card was found in the bank statement — '
                f'its charges were NOT added. Attach the right bill, or confirm the payment is in the imported period.'))

    # ---- master reconciliation ----
    draft_out = round(sum(v for (c, t, tr), v in agg.items() if t in ('Actual', 'One-off')), 2)
    draft_income = round(sum(v for (c, t, tr), v in agg.items() if t == 'Income'), 2)
    expected_out = round(recon['actual_bank'] + recon['actual_card'] + recon['card_lump'], 2)
    if abs(draft_out - expected_out) > TOL:
        issues.append(('FAIL', 'master-recon',
            f'Draft outflow {draft_out:.2f} != sum of sources {expected_out:.2f} — internal mismatch.'))
    else:
        issues.append(('INFO', 'master-recon',
            f'Draft ties out: outflow (Actual+One-off) {draft_out:.2f} (bank {recon["actual_bank"]:.2f} + '
            f'card {recon["actual_card"]:.2f} + lumps {recon["card_lump"]:.2f}); Income {draft_income:.2f}.'))
    if flagged:
        issues.append(('REVIEW', 'flags',
            f'{flagged} line item(s) need you to confirm a category (amber rows on the Bank / Card tabs).'))

    # ---- Validation tab (first) ----
    val = out.create_sheet('Validation', 0)
    val.sheet_view.showGridLines = False
    order_sev = {'FAIL': 0, 'REVIEW': 1, 'INFO': 2}
    n_fail = sum(1 for s, _, _ in issues if s == 'FAIL')
    n_rev = sum(1 for s, _, _ in issues if s == 'REVIEW')
    status = 'FAIL' if n_fail else ('REVIEW' if n_rev else 'PASS')
    t = val.cell(row=1, column=1, value=f'VALIDATION — status: {status}   '
                 f'({n_fail} blocking, {n_rev} to review)')
    t.font = Font(name=A, size=13, bold=True, color='FFFFFF')
    t.fill = FAILF if status == 'FAIL' else (PatternFill('solid', fgColor='B07D2B') if status == 'REVIEW' else OKF)
    val.cell(row=2, column=1,
             value='FAIL = do not paste yet, fix the cause. REVIEW = amber items need your category. INFO = for your record.'
             ).font = _f(9, color='808080')
    _hdr2 = ['Severity', 'Area', 'Detail']
    for j, h in enumerate(_hdr2, 1):
        c = val.cell(row=4, column=j, value=h); c.font = HF; c.fill = HFILL
    rr = 5
    for sev, area, detail in sorted(issues, key=lambda x: order_sev.get(x[0], 9)):
        val.cell(row=rr, column=1, value=sev).font = _f(9, b=(sev == 'FAIL'),
                 color=('B00000' if sev == 'FAIL' else ('B07D2B' if sev == 'REVIEW' else '777777')))
        val.cell(row=rr, column=2, value=area).font = _f(9)
        val.cell(row=rr, column=3, value=detail).font = _f(9)
        val.cell(row=rr, column=3).alignment = Alignment(wrap_text=True, vertical='top')
        rr += 1
    for col, w in zip('ABC', [10, 14, 110]): val.column_dimensions[col].width = w
    val.freeze_panes = 'A5'

    out._sheets.sort(key=lambda s: {'Validation': 0, 'Entries draft': 1, 'Bank': 2}.get(s.title, 3))
    out.save(out_path)
    return status, n_fail, n_rev

# =====================================================================
def infer_month(rows):
    ds = [t['date'] for t in rows if isinstance(t['date'], datetime.datetime)]
    if ds:
        d = min(ds); return datetime.datetime(d.year, d.month, 1)
    return datetime.datetime.today().replace(day=1)

def main():
    ap = argparse.ArgumentParser(description='Extensible, fail-loud expense categoriser.')
    ap.add_argument('bank')
    ap.add_argument('--bill', action='append', default=[])
    ap.add_argument('--month')
    ap.add_argument('--rules', default='rules.csv')
    ap.add_argument('--out', default='categorised draft.xlsx')
    a = ap.parse_args()

    issues = []
    rules = load_rules(a.rules, issues)
    fmt, rows = parse_bank(a.bank, issues)
    bills = [parse_card_bill(p, issues) for p in a.bill]
    month = datetime.datetime.strptime(a.month + '-01', '%Y-%m-%d') if a.month else infer_month(rows or [])

    if rows is None:
        # still write a Validation-only workbook so the failure is visible, not a crash
        rows = []
    status, nf, nr = build_draft(rows, bills, month, rules, issues, a.out)

    print(f'Wrote {a.out}')
    print(f'  month {month.date()}  bank rows {len(rows)}  bills {len(a.bill)}  ->  status {status} '
          f'({nf} blocking, {nr} to review)')
    for sev, area, detail in issues:
        if sev in ('FAIL', 'REVIEW'):
            print(f'  [{sev}] {area}: {detail}')
    sys.exit(2 if status == 'FAIL' else 0)

if __name__ == '__main__':
    main()
